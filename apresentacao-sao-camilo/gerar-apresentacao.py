#!/usr/bin/env python3
# Gera a apresentação do Hospital São Camilo a partir do template da marca.
# O visual (mármore, cartão branco, selo do logo) vem dos slide masters —
# só precisamos posicionar conteúdo dentro da área útil do cartão.

import copy, os
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import ImageFont

W = os.path.dirname(os.path.abspath(__file__))
FOTOS_EQ = "/Users/ericbottchergierse/Desktop/Fisioterapia Cidade Jardim/apresentacao-sao-camilo/fotos-equipe"
FOTOS_MOD = "/Users/ericbottchergierse/Desktop/Fisioterapia Cidade Jardim/apresentacao-sao-camilo/fotos-modalidades"

# ---- identidade da marca (extraída do template) ----
VERMELHO = RGBColor(0xC0, 0x00, 0x00)
VERM_CAPA = RGBColor(0xC3, 0x00, 0x0A)
CINZA = RGBColor(0x69, 0x69, 0x69)
ROSA = RGBColor(0xF7, 0xE2, 0xE2)
CINZA_CLARO = RGBColor(0xBD, 0xBD, 0xBD)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONTE = "Trebuchet MS"

# ---- área útil do cartão branco (medida no fundo do master) ----
CARD_L, CARD_R = 0.48, 12.85
CHIP_T, CHIP_H = 0.67, 0.60
TITULO_L, TITULO_T = 0.89, 0.56
CONTEUDO_T = 1.62
CONTEUDO_B = 6.15          # acima do selo do logo (canto inferior direito)

_fbold = ImageFont.truetype("/System/Library/Fonts/Supplemental/Trebuchet MS Bold.ttf", 40)


def larg_titulo(txt):
    """Largura do chip rosa, calibrada pelo template (fórmula confere com o original)."""
    pol = _fbold.getlength(txt) / 72.0
    return max(5.87, 0.69 + pol + 0.68)


def txbox(slide, l, t, w, h, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    return tb, tf


def run(par, texto, *, sz, cor, bold=False, italic=False):
    r = par.add_run()
    r.text = texto
    f = r.font
    f.name = FONTE
    f.size = Pt(sz)
    f.bold = bold
    f.italic = italic
    f.color.rgb = cor
    return r


def recuo(p, esq, primeira=None):
    """Recuo pendente: a 2ª linha em diante alinha com o texto, não com a margem.
    Sem isso, a quebra de linha de um item volta a colar na margem esquerda."""
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(esq * 914400)))
    pPr.set("indent", str(int((primeira if primeira is not None else -esq) * 914400)))


def moldura(slide, l, t, w, h, rotulo):
    """Espaço reservado para foto: moldura discreta com rótulo centralizado."""
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = RGBColor(0xF4, 0xF4, 0xF4)
    sh.line.color.rgb = CINZA_CLARO
    sh.line.width = Pt(1)
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run(p, rotulo, sz=11, cor=RGBColor(0x9E, 0x9E, 0x9E), italic=True)
    return sh


def foto_ou_moldura(slide, caminho, l, t, w, h, rotulo):
    """Insere a foto se existir; senão deixa a moldura reservada.
    A foto é encaixada por 'contain' para não distorcer."""
    if caminho and os.path.exists(caminho):
        from PIL import Image
        iw, ih = Image.open(caminho).size
        escala = min(w / iw, h / ih)
        fw, fh = iw * escala, ih * escala
        slide.shapes.add_picture(caminho, Inches(l + (w - fw) / 2), Inches(t + (h - fh) / 2),
                                 Inches(fw), Inches(fh))
        return True
    moldura(slide, l, t, w, h, rotulo)
    return False


# =====================================================================
prs = Presentation(os.path.join(W, "equipe.pptx"))
NS_R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

LAY_CAPA = prs.slide_masters[0].slide_layouts[0]
LAY_CARD = prs.slide_masters[1].slide_layouts[0]
LAY_FIM = prs.slide_masters[3].slide_layouts[0]

# Os slides descartados só são removidos DEPOIS de criar os novos: apagar
# antes faz o python-pptx reaproveitar os números de parte livres
# (ppt/slides/slide6.xml etc.) e o pacote sai com nomes duplicados.
sldIdLst = prs.slides._sldIdLst
originais = list(sldIdLst)
MANTER = {0: "capa", 5: "fim", 6: "contatos"}
s_capa, s_fim, s_contatos = prs.slides[0], prs.slides[5], prs.slides[6]
el_capa, el_fim, el_contatos = originais[0], originais[5], originais[6]


def novo_card(titulo, subtitulo=None):
    s = prs.slides.add_slide(LAY_CARD)
    chip = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.19), Inches(CHIP_T),
                              Inches(larg_titulo(titulo)), Inches(CHIP_H))
    chip.fill.solid()
    chip.fill.fore_color.rgb = ROSA
    chip.line.fill.background()
    chip.shadow.inherit = False
    chip.text_frame.paragraphs[0].add_run().text = ""

    _, tf = txbox(s, TITULO_L, TITULO_T, 11.6, 0.80, wrap=False)
    run(tf.paragraphs[0], titulo, sz=40, cor=VERMELHO, bold=True)

    if subtitulo:
        _, tf2 = txbox(s, 0.89, 1.42, 11.4, 0.42)
        run(tf2.paragraphs[0], subtitulo, sz=15, cor=CINZA, italic=True)
    return s


# =====================================================================
# 1. CAPA
for sh in s_capa.shapes:
    if sh.name == "Retângulo 2":
        p = sh.text_frame.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run(p, "Fisioterapia Cidade Jardim", sz=28, cor=VERM_CAPA, bold=True, italic=True)
    elif sh.name == "Retângulo 1":
        p = sh.text_frame.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run(p, "Nossa equipe no Hospital São Camilo Pompéia", sz=19, cor=CINZA, italic=True)

# =====================================================================
# 2. QUEM SOMOS
s = novo_card("Quem Somos", "A clínica que passa a integrar a assistência fisioterapêutica do São Camilo Pompéia")
_, tf = txbox(s, 0.89, 2.15, 6.9, 3.6)
for i, (t, d) in enumerate([
    ("Cuidado individualizado.",
     "Cada tratamento começa por uma avaliação completa e resulta em um plano próprio para o paciente."),
    ("Equipe multiespecialista.",
     "Seis fisioterapeutas com formações complementares, cobrindo do esporte à gerontologia."),
    ("Prática baseada em evidências.",
     "Condutas atualizadas, com foco em recuperação funcional e segurança assistencial."),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(24)
    p.line_spacing = 1.25
    run(p, t + " ", sz=16, cor=VERMELHO, bold=True)
    run(p, d, sz=16, cor=CINZA)

foto_ou_moldura(s, os.path.join(FOTOS_MOD, "clinica.jpg"), 8.15, 2.05, 4.2, 3.75,
                "Foto da clínica ou da equipe")

# =====================================================================
# 3. DESAFIO E OPORTUNIDADE
s = novo_card("Desafio e Oportunidade")
_, tf = txbox(s, 0.89, 1.95, 11.4, 4.0)
blocos = [
    [("Assumir a atuação em um hospital de grande porte exige ", 0), ("integração rápida", 1),
     (" aos processos assistenciais, ", 0), ("alinhamento", 1),
     (" com as equipes multiprofissionais e ", 0), ("compromisso", 1),
     (" com a qualidade e a segurança do paciente.", 0)],
    [("Ao mesmo tempo, representa uma excelente ", 0), ("oportunidade", 1),
     (" para fortalecer uma assistência fisioterapêutica especializada, baseada em evidências e focada na recuperação funcional e na melhor ", 0),
     ("experiência do paciente", 1), (".", 0)],
    [("Mais do que formar uma nova equipe, iniciamos uma jornada de colaboração, inovação e ", 0),
     ("melhoria contínua", 1), (", contribuindo para a ", 0),
     ("excelência assistencial do Hospital São Camilo", 1), (".", 0)],
]
for i, bloco in enumerate(blocos):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    p.space_after = Pt(18)
    p.line_spacing = 1.25
    for txt, destaque in bloco:
        run(p, txt, sz=16, cor=VERMELHO if destaque else CINZA, bold=bool(destaque))

# =====================================================================
# 4. NOSSOS SERVIÇOS  (ortopédica/esportiva primeiro, respiratória ambulatorial por último)
SERVICOS = [
    ("Fisioterapia Ortopédica e Esportiva", "Reabilitação musculoesquelética, pós-operatórios e retorno seguro ao esporte."),
    ("Fisioterapia Pélvica", "Disfunções do assoalho pélvico, saúde da mulher e do homem."),
    ("RPG — Reeducação Postural Global", "Correção de desequilíbrios posturais e alívio de dores crônicas."),
    ("Pilates", "Força, mobilidade, postura e consciência corporal com equipamentos."),
    ("Hidroterapia", "Reabilitação em meio aquático, com menor sobrecarga articular."),
    ("Fisioterapia Geriátrica e Gerontologia", "Funcionalidade, equilíbrio e prevenção de quedas no idoso."),
    ("Fisioterapia Neurológica", "Reabilitação funcional em doenças e lesões neurológicas."),
    ("Terapias Manuais e Liberação Miofascial", "Quiropraxia e técnicas manuais para dor e mobilidade."),
    ("Drenagem Linfática", "Redução de edemas e suporte à recuperação pós-operatória."),
    ("Fisioterapia Respiratória Ambulatorial", "Reabilitação pulmonar e suporte respiratório ambulatorial."),
]
s = novo_card("Nossos Serviços", "Uma equipe preparada para oferecer cuidado integral, humanizado e baseado em evidências")
COLS = [0.89, 6.95]
for col in range(2):
    _, tf = txbox(s, COLS[col], 2.05, 5.5, 4.0)
    for i, (nome, desc) in enumerate(SERVICOS[col * 5:col * 5 + 5]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        recuo(p, 0.42)
        run(p, f"{col * 5 + i + 1}.  ", sz=12, cor=VERMELHO, bold=True)
        run(p, nome, sz=13, cor=VERMELHO, bold=True)
        p2 = tf.add_paragraph()
        p2.space_after = Pt(14)
        recuo(p2, 0.42, primeira=0)
        run(p2, desc, sz=11, cor=CINZA)

# =====================================================================
# 5-6. MODALIDADES (espaço para fotos)
def slide_modalidades(titulo, sub, itens, base):
    s = novo_card(titulo, sub)
    n = len(itens)
    gap, esq, dir_ = 0.30, 0.89, 12.30
    larg = (dir_ - esq - gap * (n - 1)) / n
    for i, (rotulo, arq) in enumerate(itens):
        l = esq + i * (larg + gap)
        foto_ou_moldura(s, os.path.join(FOTOS_MOD, arq), l, 2.12, larg, 2.95, "Foto")
        _, tf = txbox(s, l, 5.20, larg, 0.85)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run(p, rotulo, sz=12, cor=VERMELHO, bold=True)
    return s


slide_modalidades("Modalidades", "Nossas áreas de atuação em imagens  (1 de 2)", [
    ("Ortopédica e Esportiva", "ortopedica.jpg"),
    ("Pélvica", "pelvica.jpg"),
    ("RPG", "rpg.jpg"),
    ("Pilates", "pilates.jpg"),
    ("Hidroterapia", "hidroterapia.jpg"),
], FOTOS_MOD)

slide_modalidades("Modalidades", "Nossas áreas de atuação em imagens  (2 de 2)", [
    ("Geriátrica e Gerontologia", "geriatria.jpg"),
    ("Neurológica", "neurologica.jpg"),
    ("Terapias Manuais", "manuais.jpg"),
    ("Drenagem Linfática", "drenagem.jpg"),
    ("Respiratória Ambulatorial", "respiratoria.jpg"),
], FOTOS_MOD)

# =====================================================================
# 7. VÍDEO + mini-currículo da Márcia
s = novo_card("Apresentação em Vídeo")
VW, VH = 6.40, 3.55
VL = (13.333 - VW) / 2
moldura(s, VL, 1.48, VW, VH, "Área reservada para o vídeo")
_, tf = txbox(s, 1.60, 5.22, 10.13, 0.95)
p = tf.paragraphs[0]
p.alignment = PP_ALIGN.CENTER
p.line_spacing = 1.2
run(p, "Márcia  ", sz=17, cor=VERMELHO, bold=True)
run(p, "— CREFITO 39705-F\n", sz=13, cor=CINZA, bold=True)
p2 = tf.add_paragraph()
p2.alignment = PP_ALIGN.CENTER
p2.line_spacing = 1.2
run(p2, "Fisioterapeuta e educadora física, especialista em Hidroterapia, "
        "Fisioterapia Esportiva e Geriatria e Gerontologia.", sz=13, cor=CINZA)

# =====================================================================
# 8. EQUIPE — visão geral
EQUIPE = [
    ("Márcia", "Hidroterapia · Esportiva · Gerontologia", "CREFITO 39705-F", "marcia.png",
     ["Fisioterapeuta e educadora física.",
      "Especialista em Hidroterapia, Fisioterapia Esportiva e Geriatria e Gerontologia."]),
    ("Dra. Ana Paula Volkers", "Saúde da Mulher e do Homem", "CREFITO 213821-F", "anapaula.png",
     ["Fisioterapeuta especializada em Saúde da Mulher e do Homem pela Santa Casa de São Paulo.",
      "Atuação em Fisioterapia Pélvica.",
      "Formação completa em RPG e Pilates."]),
    ("Caroline Espíndola de Andrade", "Gerontologia", "CREFITO 212383-F", "caroline.png",
     ["Fisioterapeuta pela Universidade de Mogi das Cruzes.",
      "Especializada em Gerontologia pela FMU.",
      "Formação completa em Pilates."]),
    ("Rosália Firmino da Silva", "UTI · Ortopédica · Neurológica", "CREFITO 222312", "rosalia.png",
     ["Fisioterapeuta pelo Centro Universitário Sant’Anna (2014).",
      "Pós-graduada em Fisioterapia em UTI pela Faculdade Iguaçu (2026).",
      "Formação completa em Pilates.",
      "Atuação em Fisioterapia Ortopédica, Neurológica, Respiratória e Pilates."]),
    ("Dr. Thales Romero", "Esporte e Exercício", "CREFITO 3244103-F", "thales.png",
     ["Fisioterapeuta especializado em Fisioterapia do Esporte e Exercício pelo HCFMUSP.",
      "Atuação com atletas de alto rendimento em clubes.",
      "Reabilitação e recuperação ortopédica de lesões."]),
    ("Marcelo Hirose", "Ortopédica · Fisiologia e Biomecânica", "CREFITO 115679-F", "marcelo.png",
     ["Graduado em Fisioterapia pela UNIP (2007).",
      "Pós-graduado em Fisiologia e Biomecânica pelo HCFMUSP (2009).",
      "Atuação em Fisioterapia Ortopédica.",
      "Formação completa em Pilates e RPG."]),
]

s = novo_card("Nossa Equipe", "Seis fisioterapeutas, formações complementares, um só cuidado")
gap, esq, dir_ = 0.26, 0.89, 12.30
larg = (dir_ - esq - gap * 5) / 6
for i, (nome, area, crefito, arq, _bio) in enumerate(EQUIPE):
    l = esq + i * (larg + gap)
    foto_ou_moldura(s, os.path.join(FOTOS_EQ, arq), l, 2.05, larg, 2.30, "Foto")
    _, tf = txbox(s, l - 0.05, 4.48, larg + 0.10, 1.45)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 1.05
    run(p, nome, sz=11.5, cor=VERMELHO, bold=True)
    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.line_spacing = 1.05
    p2.space_before = Pt(3)
    run(p2, area, sz=9, cor=CINZA)
    if crefito:
        p3 = tf.add_paragraph()
        p3.alignment = PP_ALIGN.CENTER
        p3.space_before = Pt(2)
        run(p3, crefito, sz=8.5, cor=CINZA, italic=True)

# =====================================================================
# 9-14. SLIDES INDIVIDUAIS
RODAPE = "Analgesia / Funcionalidade / Mobilidade / Fortalecimento / Equilíbrio / Reabilitação"
for nome, area, crefito, arq, bio in EQUIPE:
    s = novo_card("Nossa Equipe")
    foto_ou_moldura(s, os.path.join(FOTOS_EQ, arq), 1.00, 1.95, 3.10, 3.95, "Foto")

    _, tf = txbox(s, 4.45, 2.05, 7.85, 3.8)
    p = tf.paragraphs[0]
    run(p, nome, sz=26, cor=VERMELHO, bold=True)

    p = tf.add_paragraph()
    p.space_before = Pt(4)
    run(p, area + (f"   ·   {crefito}" if crefito else ""), sz=13, cor=CINZA, italic=True)

    for linha in bio:
        p = tf.add_paragraph()
        p.space_before = Pt(11)
        p.line_spacing = 1.2
        recuo(p, 0.30)
        run(p, "•  ", sz=15, cor=VERMELHO, bold=True)
        run(p, linha, sz=15, cor=CINZA)

    _, tfr = txbox(s, 1.10, 6.24, 11.0, 0.42, wrap=False)
    run(tfr.paragraphs[0], RODAPE, sz=15, cor=VERMELHO, bold=True, italic=True)

# =====================================================================
# CONTATOS — número da clínica + Instagram
for sh in s_contatos.shapes:
    if sh.name == "CaixaDeTexto 5":          # chip com ícone de arroba
        p = sh.text_frame.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run(p, "@fisioterapiacidadejardim", sz=14, cor=CINZA, bold=True, italic=True)
    elif sh.name == "CaixaDeTexto 11":       # chip com ícone de telefone
        p = sh.text_frame.paragraphs[0]
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        run(p, "(11) 98805-8005", sz=14, cor=CINZA, bold=True, italic=True)

# =====================================================================
# Agora sim: descartar os slides do template que não entram
ids = list(sldIdLst)
descartar = [originais[i] for i in range(len(originais)) if i not in MANTER]
for el in descartar:
    prs.part.drop_rel(el.get(NS_R))
    sldIdLst.remove(el)

# ORDEM FINAL: capa · conteúdo · contatos · encerramento
novos = [el for el in list(sldIdLst) if el not in (el_capa, el_fim, el_contatos)]
for el in list(sldIdLst):
    sldIdLst.remove(el)
for el in [el_capa] + novos + [el_contatos, el_fim]:
    sldIdLst.append(el)

out = "/Users/ericbottchergierse/Desktop/Fisioterapia Cidade Jardim/apresentacao-sao-camilo/Fisioterapia Cidade Jardim - Hospital Sao Camilo.pptx"
prs.save(out)
print("salvo:", out)
print("slides:", len(prs.slides.__iter__.__self__._sldIdLst))
